import sqlite3
import os
import re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


class UniversityParser:
    """
    Парсер Excel-файла с оценками студентов и загрузка их в SQLite.

    Пример использования:
        parser = UniversityParser(excel_file="file.xlsx", db_file="university.db")
        result = parser.parse()
        print(result)
    """

    # Служебные строки, которые не являются студентами
    IGNORED_FIO_VALUES = {
        "ОТЛИЧНО", "ХОРОШО", "УДОВЛЕТВ.", "НЕУДОВЛ",
        "ИТОГ", "ПРОМЕЖУТОЧНЫЙ БАЛЛ",
    }

    # Ключевые слова для определения служебных столбцов
    SERVICE_COLUMN_KEYWORDS = ["фио", "фамилия", "итог", "пропуск", "доп.задания"]

    # Названия листов в Excel
    SHEET_DATES = "Даты_темы"
    SHEET_GRADES = "Баллы, оценки"

    def __init__(self, excel_file: str, db_file: str = "university.db",
                 recreate_db: bool = True, verbose: bool = True):
        """
        :param excel_file: путь к Excel-файлу
        :param db_file: путь к SQLite-базе
        :param recreate_db: если True — удаляет старую БД перед импортом
        :param verbose: печатать ли прогресс в консоль
        """
        self.excel_file = excel_file
        self.db_file = db_file
        self.recreate_db = recreate_db
        self.verbose = verbose

        self.conn = None
        self.cursor = None
        self.wb = None

        # Словарь {номер_урока: {"дата": ..., "тема": ...}}
        self.lessons_info = {}
        # Словарь {индекс_столбца: activity_id}
        self.col_to_activity_id = {}

    # ---------------------------------------------------------------
    # Публичные методы
    # ---------------------------------------------------------------

    def parse(self) -> dict:
        """
        Основной метод: парсит Excel и загружает данные в БД.
        Возвращает словарь со статистикой.
        """
        try:
            self._log(f"Старт парсинга: {self.excel_file}")

            self._prepare_database()
            self._open_workbook()
            self._parse_lessons_info()
            self._parse_activities()
            stats = self._parse_students_and_grades()

            self._log("\nИтоговая статистика базы данных:")
            self._log(f"   Студентов: {stats['students']}")
            self._log(f"   Видов работ: {stats['activities']}")
            self._log(f"   Оценок (включая пропуски): {stats['grades']}")
            self._log(f"\nГотово! База данных: {self.db_file}")

            return stats

        finally:
            self._cleanup()

    # ---------------------------------------------------------------
    # Подготовка
    # ---------------------------------------------------------------

    def _log(self, message: str):
        if self.verbose:
            print(message)

    def _prepare_database(self):
        """Создаёт (или пересоздаёт) SQLite-базу и таблицы."""
        if self.recreate_db and os.path.exists(self.db_file):
            try:
                os.remove(self.db_file)
                self._log(f"Старая база {self.db_file} удалена.")
            except PermissionError:
                raise PermissionError(
                    f"Файл {self.db_file} занят другой программой. "
                    f"Закройте DB Browser for SQLite и попробуйте снова."
                )

        self.conn = sqlite3.connect(self.db_file)
        self.cursor = self.conn.cursor()
        self._create_tables()
        self.conn.commit()
        self._log("Таблицы в БД созданы.")

    def _create_tables(self):
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS Students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                fio TEXT NOT NULL,
                itog_ball REAL,
                kol_propuskov INTEGER
            )
        ''')
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS Activities (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nazvanie TEXT NOT NULL,
                tip TEXT NOT NULL,
                data TEXT,
                tema TEXT,
                max_ball REAL
            )
        ''')
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS Grades (
                student_id INTEGER,
                activity_id INTEGER,
                poluchenny_ball REAL,
                PRIMARY KEY (student_id, activity_id),
                FOREIGN KEY (student_id) REFERENCES Students(id),
                FOREIGN KEY (activity_id) REFERENCES Activities(id)
            )
        ''')

    def _open_workbook(self):
        self.wb = load_workbook(self.excel_file, data_only=True)

    # ---------------------------------------------------------------
    # Парсинг листа "Даты_темы"
    # ---------------------------------------------------------------

    def _parse_lessons_info(self):
        self._log(f"\nЧитаем лист '{self.SHEET_DATES}'...")
        sheet = self.wb[self.SHEET_DATES]

        # Первый проход: собираем строки, где явно указан номер урока
        for row in sheet.iter_rows(min_row=5, values_only=True):
            lesson_num_raw = row[1]
            if lesson_num_raw is None:
                continue
            try:
                lesson_num = int(lesson_num_raw)
            except (ValueError, TypeError):
                continue

            date_val = row[2] if row[2] else ""
            date_str = str(date_val).split(" ")[0] if date_val else ""
            tema_part = str(row[3]).strip() if row[3] else ""

            if lesson_num not in self.lessons_info:
                self.lessons_info[lesson_num] = {"дата": date_str, "тема": tema_part}
            else:
                if tema_part and tema_part not in self.lessons_info[lesson_num]["тема"]:
                    self.lessons_info[lesson_num]["тема"] += " " + tema_part

        # Второй проход: подтягиваем продолжения тем из строк без номера урока
        for i in range(5, sheet.max_row + 1):
            lesson_cell = sheet.cell(row=i, column=2).value
            tema_cell = sheet.cell(row=i, column=4).value

            if lesson_cell:
                try:
                    lesson_num = int(lesson_cell)
                    if tema_cell and lesson_num in self.lessons_info:
                        current = self.lessons_info[lesson_num]["тема"]
                        new_part = str(tema_cell).strip()
                        if new_part and new_part not in current:
                            self.lessons_info[lesson_num]["тема"] += " " + new_part
                except (ValueError, TypeError):
                    pass
            elif tema_cell and self.lessons_info:
                last_lesson = max(self.lessons_info.keys())
                new_part = str(tema_cell).strip()
                if new_part and new_part not in self.lessons_info[last_lesson]["тема"]:
                    self.lessons_info[last_lesson]["тема"] += " " + new_part

        self._log(f"   Найдено {len(self.lessons_info)} уроков с темами.")

    # ---------------------------------------------------------------
    # Парсинг листа "Баллы, оценки" → Activities
    # ---------------------------------------------------------------

    def _parse_activities(self):
        self._log(f"\nЧитаем лист '{self.SHEET_GRADES}'...")
        sheet = self.wb[self.SHEET_GRADES]

        header_row1 = [cell.value for cell in sheet[1]]
        header_row2 = [cell.value for cell in sheet[2]]

        activity_columns = self._collect_activity_columns(header_row1, header_row2)
        self._log(f"   Найдено {len(activity_columns)} видов работ для импорта.")

        for act in activity_columns:
            activity_id = self._insert_activity(act)
            self.col_to_activity_id[act["col_idx"]] = activity_id

        self.conn.commit()
        self._log(f"В таблицу Activities добавлено {len(activity_columns)} записей.")

    def _collect_activity_columns(self, header_row1, header_row2) -> list:
        """Возвращает список словарей с описанием столбцов-активностей."""
        columns = []
        for col_idx in range(len(header_row1)):
            name = header_row1[col_idx]
            date_val = header_row2[col_idx]

            if name is None:
                continue

            name_str = str(name).strip()
            name_lower = name_str.lower()

            if any(kw in name_lower for kw in self.SERVICE_COLUMN_KEYWORDS):
                continue
            if not name_str and not date_val:
                continue

            date_str = str(date_val).split(" ")[0] if date_val else ""

            columns.append({
                "col_idx": col_idx,
                "name": name_str,
                "date": date_str,
            })
        return columns

    def _insert_activity(self, act: dict) -> int:
        """Определяет тип, тему, max_ball и вставляет запись в Activities."""
        name = act["name"]
        date = act["date"]
        name_lower = name.lower()

        # Определяем тип и максимальный балл
        if name_lower.startswith("пр"):
            tip, max_ball = "ПР", 5
        elif name_lower.startswith("лр"):
            tip, max_ball = "ЛР", 5
        elif "оценка" in name_lower or "зачет" in name_lower:
            tip, max_ball = "Зачет", 10
        elif "дз" in name_lower:
            tip, max_ball = "ДЗ", 3
        elif re.match(r'^\d+$', name):
            tip, max_ball = "Урок", 2
        else:
            tip, max_ball = "Урок", 2

        # Пытаемся найти тему
        tema = ""
        match = re.search(r'(?:ур\.?|урок)\s*(\d+)', name, re.IGNORECASE)
        if not match:
            match = re.search(r'^(\d+)$', name)

        if match:
            lesson_num = int(match.group(1))
            if lesson_num in self.lessons_info:
                tema = self.lessons_info[lesson_num]["тема"]
                if not date:
                    date = self.lessons_info[lesson_num]["дата"]

        if not tema and date:
            for _, l_data in self.lessons_info.items():
                if l_data["дата"] == date:
                    tema = l_data["тема"]
                    break

        self.cursor.execute('''
            INSERT INTO Activities (nazvanie, tip, data, tema, max_ball)
            VALUES (?, ?, ?, ?, ?)
        ''', (name, tip, date, tema, max_ball))

        return self.cursor.lastrowid

    # ---------------------------------------------------------------
    # Парсинг студентов и оценок
    # ---------------------------------------------------------------

    def _parse_students_and_grades(self) -> dict:
        self._log("\nЧитаем студентов...")
        sheet = self.wb[self.SHEET_GRADES]
        header_row1 = [cell.value for cell in sheet[1]]

        fio_col, itog_col, propusk_col = self._find_service_columns(header_row1)

        if fio_col is None:
            raise ValueError("Не найден столбец с ФИО в листе 'Баллы, оценки'.")

        self._log(f"   Столбец ФИО: {get_column_letter(fio_col + 1)}")
        self._log(f"   Столбец Итог: "
                  f"{get_column_letter(itog_col + 1) if itog_col is not None else 'Не найден'}")
        self._log(f"   Столбец Пропуски: "
                  f"{get_column_letter(propusk_col + 1) if propusk_col is not None else 'Не найден'}")

        students_count = 0
        grades_count = 0

        for row in sheet.iter_rows(min_row=3, values_only=True):
            student_id = self._process_student_row(
                row, fio_col, itog_col, propusk_col
            )
            if student_id is None:
                continue

            students_count += 1
            grades_count += self._insert_grades(student_id, row)

        self.conn.commit()
        self._log(f"Добавлено {students_count} студентов и {grades_count} оценок.")

        return {
            "students": students_count,
            "activities": len(self.col_to_activity_id),
            "grades": grades_count,
        }

    def _find_service_columns(self, header_row1):
        """Находит индексы столбцов ФИО, Итог, Пропуски."""
        fio_col = itog_col = propusk_col = None
        for col_idx, val in enumerate(header_row1):
            if val is None:
                continue
            val_str = str(val).strip().lower()
            if "фио" in val_str or "фамилия" in val_str:
                fio_col = col_idx
            elif "итог" in val_str:
                itog_col = col_idx
            elif "пропуск" in val_str:
                propusk_col = col_idx
        return fio_col, itog_col, propusk_col

    def _process_student_row(self, row, fio_col, itog_col, propusk_col):
        """
        Обрабатывает одну строку со студентом.
        Возвращает student_id или None, если строка не является студентом.
        """
        if len(row) <= fio_col:
            return None

        fio_val = row[fio_col]
        if fio_val is None or str(fio_val).strip() == "":
            return None

        fio = str(fio_val).strip()

        if fio.upper() in self.IGNORED_FIO_VALUES:
            return None

        itog = self._parse_float(row, itog_col)
        propuski = self._parse_int(row, propusk_col)

        # Пропускаем служебные строки (нет ни итога, ни пропусков)
        if itog is None and propuski is None:
            return None

        self.cursor.execute('''
            INSERT INTO Students (fio, itog_ball, kol_propuskov)
            VALUES (?, ?, ?)
        ''', (fio, itog, propuski))
        return self.cursor.lastrowid

    def _insert_grades(self, student_id: int, row) -> int:
        """Вставляет оценки студента по всем активностям. Возвращает их количество."""
        count = 0
        for col_idx, activity_id in self.col_to_activity_id.items():
            if col_idx >= len(row):
                continue

            ball = self._parse_float(row, col_idx, skip_formulas=True)

            self.cursor.execute('''
                INSERT INTO Grades (student_id, activity_id, poluchenny_ball)
                VALUES (?, ?, ?)
            ''', (student_id, activity_id, ball))
            count += 1
        return count

    # ---------------------------------------------------------------
    # Утилиты
    # ---------------------------------------------------------------

    @staticmethod
    def _parse_float(row, col_idx, skip_formulas: bool = False):
        """Извлекает float из ячейки. Возвращает None, если не получилось."""
        if col_idx is None or col_idx >= len(row):
            return None
        val = row[col_idx]
        if val is None:
            return None
        if isinstance(val, (int, float)):
            return float(val)
        if isinstance(val, str):
            if skip_formulas and val.startswith("="):
                return None
            try:
                return float(val.replace(",", "."))
            except ValueError:
                return None
        return None

    @staticmethod
    def _parse_int(row, col_idx):
        """Извлекает int из ячейки. Возвращает None, если не получилось."""
        if col_idx is None or col_idx >= len(row):
            return None
        val = row[col_idx]
        if val is None:
            return None
        if isinstance(val, (int, float)):
            return int(val)
        if isinstance(val, str) and val.strip().isdigit():
            return int(val)
        return None

    def _cleanup(self):
        """Закрывает соединения с БД и Excel."""
        if self.conn:
            self.conn.close()
        if self.wb:
            self.wb.close()